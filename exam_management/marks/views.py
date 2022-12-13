from django.shortcuts import render,HttpResponse
from adminpage.models import *
from openpyxl import Workbook
import openpyxl
import mysql.connector
import json
import ast

config={
'host':"localhost",
'user':"root",
'password':"",
'database':"python"
}

def markshome(request):
    d=Subject.objects.all().values()
    details={'subjectlist':d}
    return render(request,"marks.html",details)

def adminpage(request):
    return render(request,"admin.html",{})

def marksoption(request):
    data=request.POST
    d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
    details={'subject':d[0],'year':data['year']}
    return render(request,"marksoption.html",details)

# ====================  THEORY ===========================

def theory(request):    
    data=request.POST
    d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
    g=Grade.objects.all().values()

    mydb=mysql.connector.connect(**config)
    mycursor = mydb.cursor()
    mycursor.execute("SELECT DISTINCT SUBSTRING(enrollment, 1, 2) AS batch,sem as sem FROM adminpage_studentdetails WHERE institute_code='"+d[0]['institute_code']+"' and program_code='"+d[0]['program_code']+"';")
    myresult = mycursor.fetchall()
    mydb.close()

    details={
        'subject':d[0],
        'year':data['year'],
        'type':data['type'],
        'grade':g,
        'batchlist':myresult,
    }
    return render(request,"theory.html",details)


def uploadtheory(request):
    data=request.POST

    if(data['btype']=='remedial'):
        mydb=mysql.connector.connect(**config)
        d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
        myresult=[]
        myresult.clear()
        d=d[0]
        b='{"Status": "F"%'
        sub=str(d['sem'])+"_"+d['subjectcode']+data['type']
        # c=('SELECT  adminpage_studentmarks.enrollment,adminpage_studentdetails.name FROM adminpage_studentmarks INNER JOIN adminpage_studentdetails ON adminpage_studentdetails.enrollment=adminpage_studentmarks.enrollment WHERE '+str(sub)+' LIKE '+"'"+b+"' AND adminpage_studentmarks.enrollment LIKE '"+data['batch']+"%'" )
        c=('''
            SELECT  adminpage_studentmarks.enrollment,adminpage_studentdetails.name,
            adminpage_studentmarks.'''+sub+''' FROM adminpage_studentmarks INNER JOIN adminpage_studentdetails ON adminpage_studentdetails.enrollment=adminpage_studentmarks.enrollment WHERE 
            '''+sub+''' LIKE '''+"'"+str(b)+"'"+''' AND 
            adminpage_studentmarks.enrollment LIKE '''+"'"+str(data['batch']+"%"+"'"))
        print(c)
        mycursor = mydb.cursor()
        mycursor.execute(c)
        myresult = mycursor.fetchall()
        mydb.close()

        arr = list()
        for i in range(0, len(myresult)):
            if(myresult[i][2]=='n'):
                arr.insert(i,[myresult[i][0],myresult[i][1],""])
            else:
                marks = ast.literal_eval(myresult[i][2])
                if str(int(data['year'])-1) not in marks['year'].keys():return HttpResponse("No Data Found");break
                print(marks['year'][str(int(data['year'])-1)]['marks'])

                if data['year'] in marks['year'].keys():
                    arr.insert(i,[(myresult[i][0]),myresult[i][1],marks['year'][data['year']]['marks'],marks['year'][str(int(data['year'])-1)]['marks']])
                else:arr.insert(i,[myresult[i][0],myresult[i][1],"000",marks['year'][str(int(data['year'])-1)]['marks']])

        d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
        g=Grade.objects.all().values()
        passing=g[len(g)-1]['r2']

        detail={'excel':arr,'subject':d[0],'grade':g,'year':data['year'],'passing':passing,'remedial':True,'remedialyear':str(int(data['year'])-1)}

        return render(request,"theoryexcel.html",detail)
    
    else:
        mydb=mysql.connector.connect(**config)
        d=Subject.objects.filter(subjectcode=data['subjectcode']).values()

        d=d[0]
        sub=str(d['sem'])+"_"+d['subjectcode']+data['type']
        c=('''SELECT  adminpage_studentmarks.enrollment,adminpage_studentdetails.name,
        adminpage_studentmarks.'''+sub+''' FROM adminpage_studentmarks INNER JOIN adminpage_studentdetails ON adminpage_studentdetails.enrollment=adminpage_studentmarks.enrollment WHERE 
        adminpage_studentdetails.sem='''+str(d['sem'])+''' AND 
        adminpage_studentdetails.institute_code='''+str(d['institute_code'])+''' AND 
        adminpage_studentdetails.program_code='''+str(d['program_code'])+''';''')

        print(c)
        mycursor = mydb.cursor()
        mycursor.execute(c)
        myresult = mycursor.fetchall()
        mydb.close()

        arr = list()
        for i in range(0, len(myresult)):
            if(myresult[i][2]=='n'):
                arr.insert(i,[myresult[i][0],myresult[i][1],""])
            else:
                marks = ast.literal_eval(myresult[i][2])
                if data['year'] in marks['year'].keys():
                    # arr[i][2]=
                    arr.insert(i,[(myresult[i][0]),myresult[i][1],marks['year'][data['year']]['marks']])
                else:arr.insert(i,[myresult[i][0],myresult[i][1],""])

        d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
        g=Grade.objects.all().values()
        passing=g[len(g)-1]['r2']

        detail={'excel':arr,'subject':d[0],'grade':g,'year':data['year'],'passing':passing}

        return render(request,"theoryexcel.html",detail)



def submittheory(request):
    data=request.POST
    data=dict(data)

    d=Subject.objects.filter(subjectcode=data['subjectcode'][0]).values()
    d=d[0]
    sub=str(d['sem'])+"_"+d['subjectcode']+"_t"
    
    
    marks={}
    for d in range(0,len(data['enrollment'])):
        mydb=mysql.connector.connect(**config)
        mycursor = mydb.cursor()
        mycursor.execute("SELECT "+sub+" FROM adminpage_studentmarks WHERE enrollment='"+data['enrollment'][d]+"'")
        myresult = mycursor.fetchall()

        if(myresult[0][0]=="n"):
            mydb=mysql.connector.connect(**config)
            mycursor = mydb.cursor()
            marks={
                'Status':data['status'][d],
                'Grade':data['grade'][d],
                'year':{
                    str(data['year'][0]):{data['type'][0]:{'marks':data['marks'][d],'out':data['total'][0],'grade':data['grade'][d], 'credit':data['gradepoint'][d]},
                    },
                }
            }
            res = json.dumps(marks)

            c=("INSERT INTO adminpage_studentmarks ("+sub+") VALUES('"+str(marks)+"') WHERE enrollment='"+data['enrollment'][d]+"'")
            c="UPDATE `adminpage_studentmarks` SET `"+sub+"` = '"+res+"' WHERE `adminpage_studentmarks`.`enrollment` = '"+data['enrollment'][d]+"';"
            mycursor.execute(c)
            mydb.commit()
            mydb.close()
            print(mycursor.rowcount, "record inserted.")

        else:
            mydb=mysql.connector.connect(**config)
            mycursor = mydb.cursor()
            marks = ast.literal_eval(myresult[0][0])
            print(marks)
            marks['Status']=data['status'][d]
            marks['Grade']=data['grade'][d]

            add={'marks':data['marks'][d],'out':data['total'][0],'grade':data['grade'][d], 'credit':data['gradepoint'][d]}
            
            marks['year'][data['year'][0]]=add
            
            res = json.dumps(marks)

            c=("INSERT INTO adminpage_studentmarks ("+sub+") VALUES('"+str(marks)+"') WHERE enrollment='"+data['enrollment'][d]+"'")
            c="UPDATE `adminpage_studentmarks` SET `"+sub+"` = '"+res+"' WHERE `adminpage_studentmarks`.`enrollment` = '"+data['enrollment'][d]+"';"
            mycursor.execute(c)
            mydb.commit()
            mydb.close()
            print(mycursor.rowcount, "record inserted.")

    return HttpResponse("MARKS ADDED SUCCESSFULLY")


# =========================  PRACTICAL =======================================================

def practical(request):
    data=request.POST
    d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
    g=Grade.objects.all().values()
    details={
        'subject':d[0],'year':data['year'],'type':data['type'],'grade':g
    }
    return render(request,"practical.html",details)


def uploadpractical(request):

    if "GET" == request.method:
        return render(request, 'myapp/index.html', {})
    else:
        data=request.POST
        excel_file = request.FILES["practicalexcel"]
        d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
        g=Grade.objects.all().values()
        wb = openpyxl.load_workbook(excel_file)

        worksheet = wb["Students Seat"]
        excel_data = list()

        for row in worksheet.iter_rows():
            row_data = list()
            for cell in row:
                row_data.append(str(cell.value))
            excel_data.append(row_data)
        passing=g[len(g)-1]['r2']

        detail={'excel':excel_data,'subject':d[0],'grade':g,'year':data['year'],'passing':passing}

        return render(request,"practicalexcel.html",detail)


def submitpractical(request):
    data=request.POST
    data=dict(data)

    d=Subject.objects.filter(subjectcode=data['subjectcode'][0]).values()
    d=d[0]
    sub=str(d['sem'])+"_"+d['subjectcode']+"_p"
    
    
    marks={}
    for d in range(0,len(data['enrollment'])):
        mydb=mysql.connector.connect(**config)
        mycursor = mydb.cursor()
        mycursor.execute("SELECT "+sub+" FROM adminpage_studentmarks WHERE enrollment='"+data['enrollment'][d]+"'")
        myresult = mycursor.fetchall()

        if(myresult[0][0]=="n"):
            mydb=mysql.connector.connect(**config)
            mycursor = mydb.cursor()
            marks={
                'Status':data['status'][d],
                'Grade':data['grade'][d],
                'year':{
                    str(data['year'][0]):{data['type'][0]:{'marks':data['marks'][d],'out':data['total'][0],'grade':data['grade'][d], 'credit':data['gradepoint'][d]},
                    },
                }
            }
            res = json.dumps(marks)

            c=("INSERT INTO adminpage_studentmarks ("+sub+") VALUES('"+str(marks)+"') WHERE enrollment='"+data['enrollment'][d]+"'")
            c="UPDATE `adminpage_studentmarks` SET `"+sub+"` = '"+res+"' WHERE `adminpage_studentmarks`.`enrollment` = '"+data['enrollment'][d]+"';"
            mycursor.execute(c)
            mydb.commit()
            mydb.close()
            print(mycursor.rowcount, "record inserted.")

        else:
            mydb=mysql.connector.connect(**config)
            mycursor = mydb.cursor()
            marks = ast.literal_eval(myresult[0][0])
            print(marks)
            marks['Status']=data['status'][d]
            marks['Grade']=data['grade'][d]

            add={'marks':data['marks'][d],'out':data['total'][0],'grade':data['grade'][d], 'credit':data['gradepoint'][d]}
            
            marks['year'][data['year'][0]]=add
            
            res = json.dumps(marks)

            c=("INSERT INTO adminpage_studentmarks ("+sub+") VALUES('"+str(marks)+"') WHERE enrollment='"+data['enrollment'][d]+"'")
            c="UPDATE `adminpage_studentmarks` SET `"+sub+"` = '"+res+"' WHERE `adminpage_studentmarks`.`enrollment` = '"+data['enrollment'][d]+"';"
            mycursor.execute(c)
            mydb.commit()
            mydb.close()
            print(mycursor.rowcount, "record inserted.")

    return HttpResponse("MARKS ADDED SUCCESSFULLY")


# =============================  MID =============================




def mid(request):
    data=request.POST
    d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
    g=Grade.objects.all().values()
    details={
        'subject':d[0],'year':data['year'],'type':data['type'],'grade':g
    }
    return render(request,"mid.html",details)


def uploadmid(request):

    if "GET" == request.method:
        return render(request, 'myapp/index.html', {})
    else:
        data=request.POST
        excel_file = request.FILES["midexcel"]
        d=Subject.objects.filter(subjectcode=data['subjectcode']).values()
        g=Grade.objects.all().values()
        wb = openpyxl.load_workbook(excel_file)

        worksheet = wb["Students Seat"]
        excel_data = list()

        for row in worksheet.iter_rows():
            row_data = list()
            for cell in row:
                row_data.append(str(cell.value))
            excel_data.append(row_data)
        passing=g[len(g)-1]['r2']

        detail={'excel':excel_data,'subject':d[0],'grade':g,'year':data['year'],'passing':passing}

        return render(request,"midexcel.html",detail)


def submitmid(request):
    data=request.POST
    data=dict(data)

    d=Subject.objects.filter(subjectcode=data['subjectcode'][0]).values()
    d=d[0]
    sub=str(d['sem'])+"_"+d['subjectcode']+"_m"
    
    
    marks={}
    for d in range(0,len(data['enrollment'])):
        mydb=mysql.connector.connect(**config)
        mycursor = mydb.cursor()
        mycursor.execute("SELECT "+sub+" FROM adminpage_studentmarks WHERE enrollment='"+data['enrollment'][d]+"'")
        myresult = mycursor.fetchall()

        if(myresult[0][0]=="n"):
            mydb=mysql.connector.connect(**config)
            mycursor = mydb.cursor()
            marks={
                'Status':data['status'][d],
                'Grade':data['grade'][d],
                'year':{
                    str(data['year'][0]):{data['type'][0]:{'marks':data['marks'][d],'out':data['total'][0],'grade':data['grade'][d], 'credit':data['gradepoint'][d]},
                    },
                }
            }
            res = json.dumps(marks)

            c=("INSERT INTO adminpage_studentmarks ("+sub+") VALUES('"+str(marks)+"') WHERE enrollment='"+data['enrollment'][d]+"'")
            c="UPDATE `adminpage_studentmarks` SET `"+sub+"` = '"+res+"' WHERE `adminpage_studentmarks`.`enrollment` = '"+data['enrollment'][d]+"';"
            mycursor.execute(c)
            mydb.commit()
            mydb.close()
            print(mycursor.rowcount, "record inserted.")

        else:
            mydb=mysql.connector.connect(**config)
            mycursor = mydb.cursor()
            marks = ast.literal_eval(myresult[0][0])
            print(marks)
            marks['Status']=data['status'][d]
            marks['Grade']=data['grade'][d]

            add={'marks':data['marks'][d],'out':data['total'][0],'grade':data['grade'][d], 'credit':data['gradepoint'][d]}
            
            marks['year'][data['year'][0]]=add
            
            res = json.dumps(marks)

            c=("INSERT INTO adminpage_studentmarks ("+sub+") VALUES('"+str(marks)+"') WHERE enrollment='"+data['enrollment'][d]+"'")
            c="UPDATE `adminpage_studentmarks` SET `"+sub+"` = '"+res+"' WHERE `adminpage_studentmarks`.`enrollment` = '"+data['enrollment'][d]+"';"
            mycursor.execute(c)
            mydb.commit()
            mydb.close()
            print(mycursor.rowcount, "record inserted.")

    return HttpResponse("MARKS ADDED SUCCESSFULLY")



